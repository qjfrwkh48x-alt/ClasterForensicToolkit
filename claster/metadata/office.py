"""
Metadata extraction from Microsoft Office documents (DOCX, XLSX, PPTX).
"""

from pathlib import Path
from typing import Dict, Any
from datetime import datetime

from claster.core.logger import get_logger

logger = get_logger(__name__)


def get_office_metadata(docx_path: str) -> Dict[str, Any]:
    """
    Extract metadata from a Word document (.docx).

    Args:
        docx_path: Path to .docx file.

    Returns:
        Dictionary with author, title, created/modified dates, etc.
    """
    try:
        from docx import Document
    except ImportError:
        logger.error("python-docx is required for DOCX metadata.")
        return {}

    path = Path(docx_path)
    if not path.exists():
        raise FileNotFoundError(f"File not found: {path}")

    try:
        doc = Document(path)
        props = doc.core_properties
        return {
            'author': props.author,
            'title': props.title,
            'subject': props.subject,
            'keywords': props.keywords,
            'created': props.created.isoformat() if props.created else None,
            'modified': props.modified.isoformat() if props.modified else None,
            'last_modified_by': props.last_modified_by,
            'revision': props.revision,
            'category': props.category,
            'comments': props.comments,
        }
    except Exception as e:
        logger.error(f"Failed to read DOCX metadata: {e}")
        return {}


def get_excel_metadata(xlsx_path: str) -> Dict[str, Any]:
    """
    Extract metadata from an Excel workbook (.xlsx).

    Args:
        xlsx_path: Path to .xlsx file.

    Returns:
        Dictionary with author, title, created/modified, etc.
    """
    try:
        from openpyxl import load_workbook
    except ImportError:
        logger.error("openpyxl is required for XLSX metadata.")
        return {}

    path = Path(xlsx_path)
    if not path.exists():
        raise FileNotFoundError(f"File not found: {path}")

    try:
        wb = load_workbook(path, read_only=True)
        props = wb.properties
        return {
            'creator': props.creator,
            'title': props.title,
            'subject': props.subject,
            'description': props.description,
            'keywords': props.keywords,
            'created': props.created.isoformat() if props.created else None,
            'modified': props.modified.isoformat() if props.modified else None,
            'last_modified_by': props.lastModifiedBy,
            'category': props.category,
        }
    except Exception as e:
        logger.error(f"Failed to read XLSX metadata: {e}")
        return {}


def get_ppt_metadata(pptx_path: str) -> Dict[str, Any]:
    """
    Extract metadata from a PowerPoint presentation (.pptx).

    Args:
        pptx_path: Path to .pptx file.

    Returns:
        Dictionary with author, title, created/modified, etc.
    """
    try:
        from pptx import Presentation
    except ImportError:
        logger.error("python-pptx is required for PPTX metadata.")
        return {}

    path = Path(pptx_path)
    if not path.exists():
        raise FileNotFoundError(f"File not found: {path}")

    try:
        prs = Presentation(path)
        props = prs.core_properties
        return {
            'author': props.author,
            'title': props.title,
            'subject': props.subject,
            'keywords': props.keywords,
            'created': props.created.isoformat() if props.created else None,
            'modified': props.modified.isoformat() if props.modified else None,
            'last_modified_by': props.last_modified_by,
            'category': props.category,
            'comments': props.comments,
        }
    except Exception as e:
        logger.error(f"Failed to read PPTX metadata: {e}")
        return {}